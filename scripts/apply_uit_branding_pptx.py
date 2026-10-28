"""
Script to apply UIT branding (Logo + Header + Footer) to PowerPoint slides.
Header (Top Left):
  - Logo: report/img/logo_uit.png
  - Line 1: ĐẠI HỌC QUỐC GIA TP. HỒ CHÍ MINH
  - Line 2: TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN
Footer (Bottom Left - All Slides):
  - Trường Đại học Công nghệ Thông tin · ĐHQG-TPHCM
"""
import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

LOGO_PATH = os.path.join('report', 'img', 'logo_uit.png')

# UIT Brand Colors
COLOR_SLATE = RGBColor(100, 116, 139)     # Gray for ĐHQG
COLOR_UIT_BLUE = RGBColor(27, 79, 154)    # Official UIT Blue
COLOR_FOOTER = RGBColor(148, 163, 184)    # Subtle gray footer

def add_uit_branding(slide, slide_width, slide_height, is_title_slide=False):
    """Adds the UIT logo, header text, and footer text to a slide."""
    
    # 1. Header: Logo UIT
    if os.path.exists(LOGO_PATH):
        logo_left = Inches(0.5)
        logo_top = Inches(0.35)
        logo_height = Inches(0.58)
        slide.shapes.add_picture(LOGO_PATH, logo_left, logo_top, height=logo_height)
    
    # 2. Header Text next to Logo
    header_left = Inches(1.15)
    header_top = Inches(0.32)
    header_width = Inches(4.5)
    header_height = Inches(0.65)
    
    txBox = slide.shapes.add_textbox(header_left, header_top, header_width, header_height)
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    # Line 1: ĐẠI HỌC QUỐC GIA TP. HỒ CHÍ MINH
    p1 = tf.paragraphs[0]
    p1.text = "ĐẠI HỌC QUỐC GIA TP. HỒ CHÍ MINH"
    p1.font.name = "Segoe UI"
    p1.font.size = Pt(8.5)
    p1.font.bold = False
    p1.font.color.rgb = COLOR_SLATE
    
    # Line 2: TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN
    p2 = tf.add_paragraph()
    p2.text = "TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN"
    p2.font.name = "Segoe UI"
    p2.font.size = Pt(11)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_UIT_BLUE
    
    # 3. Footer (Bottom Left - All slides)
    footer_left = Inches(0.5)
    footer_top = slide_height - Inches(0.45)
    footer_width = Inches(6.0)
    footer_height = Inches(0.3)
    
    ftBox = slide.shapes.add_textbox(footer_left, footer_top, footer_width, footer_height)
    ft_tf = ftBox.text_frame
    ft_tf.word_wrap = True
    ft_tf.margin_left = ft_tf.margin_top = ft_tf.margin_right = ft_tf.margin_bottom = 0
    
    fp = ft_tf.paragraphs[0]
    fp.text = "Trường Đại học Công nghệ Thông tin · ĐHQG-TPHCM"
    fp.font.name = "Segoe UI"
    fp.font.size = Pt(8.5)
    fp.font.color.rgb = COLOR_FOOTER

def apply_branding_to_presentation(input_pptx_path, output_pptx_path=None):
    """Applies UIT header and footer to all slides in an existing presentation."""
    if not os.path.exists(input_pptx_path):
        print(f"Error: File {input_pptx_path} not found.")
        return False
    
    if output_pptx_path is None:
        base, ext = os.path.splitext(input_pptx_path)
        output_pptx_path = f"{base}_uit_branded{ext}"
        
    prs = Presentation(input_pptx_path)
    slide_w = prs.slide_width
    slide_h = prs.slide_height
    
    print(f"Applying UIT branding to {len(prs.slides)} slides in {input_pptx_path}...")
    for idx, slide in enumerate(prs.slides):
        add_uit_branding(slide, slide_w, slide_h, is_title_slide=(idx == 0))
        
    prs.save(output_pptx_path)
    print(f"Saved branded presentation to: {output_pptx_path}")
    return True

if __name__ == "__main__":
    if len(sys.argv) > 1:
        apply_branding_to_presentation(sys.argv[1])
    else:
        print("Usage: python scripts/apply_uit_branding_pptx.py <path_to_presentation.pptx>")
