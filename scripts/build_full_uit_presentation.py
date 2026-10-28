"""
Script to build the COMPLETE 23-slide PowerPoint presentation with full UIT branding.
Layout and styles:
- 16:9 Widescreen (13.333 x 7.5 inches)
- Top Left: UIT Logo + ĐẠI HỌC QUỐC GIA TP. HỒ CHÍ MINH / TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN
- Bottom Left: Trường Đại học Công nghệ Thông tin · ĐHQG-TPHCM (on ALL 23 slides)
- Slide 01: Formatted matching the user's reference image (cards, author list, monitor mockup)
- Slides 02-23: Ultra-concise keyword-first bullet points, bold key metrics, high-res charts, embedded GIF
- Presenter Notes: Complete with Transition, Element Animations, Media references, and Speaker Script
"""
import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

if sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

OUTPUT_PPTX = os.path.join('report', 'ChiliFarm_UIT_Presentation.pptx')
LOGO_PATH = os.path.join('report', 'img', 'logo_uit.png')
IMG_DIR = os.path.join('report', 'img')

# Color Palette
COLOR_BG_WHITE = RGBColor(255, 255, 255)
COLOR_BG_LIGHT = RGBColor(248, 250, 252)    # Slate 50
COLOR_TEXT_MAIN = RGBColor(15, 23, 42)      # Slate 900
COLOR_TEXT_MUTED = RGBColor(100, 116, 139)  # Slate 500
COLOR_TEXT_FOOTER = RGBColor(148, 163, 184) # Slate 400
COLOR_UIT_BLUE = RGBColor(27, 79, 154)      # Official UIT Blue (#1b4f9a)
COLOR_NAVY_DARK = RGBColor(15, 30, 60)      # Navy Dark
COLOR_CARD_BLUE = RGBColor(238, 246, 255)   # Light blue tint card
COLOR_ORANGE_TAG = RGBColor(234, 88, 12)    # Orange for "BÁO CÁO ĐỀ TÀI"
COLOR_ACCENT_TEAL = RGBColor(13, 148, 136)  # Teal accent

def add_branding(slide):
    # 1. Header Logo
    if os.path.exists(LOGO_PATH):
        slide.shapes.add_picture(LOGO_PATH, Inches(0.6), Inches(0.35), height=Inches(0.65))
    
    # 2. Header Text
    txBox = slide.shapes.add_textbox(Inches(1.35), Inches(0.35), Inches(5.5), Inches(0.7))
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p1 = tf.paragraphs[0]
    p1.text = "ĐẠI HỌC QUỐC GIA TP. HỒ CHÍ MINH"
    p1.font.name = "Segoe UI"
    p1.font.size = Pt(8.5)
    p1.font.bold = False
    p1.font.color.rgb = COLOR_TEXT_MUTED
    
    p2 = tf.add_paragraph()
    p2.text = "TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN"
    p2.font.name = "Segoe UI"
    p2.font.size = Pt(11.5)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_UIT_BLUE

    # 3. Footer (Bottom Left)
    ftBox = slide.shapes.add_textbox(Inches(0.6), Inches(7.05), Inches(6.0), Inches(0.35))
    ft_tf = ftBox.text_frame
    ft_tf.word_wrap = True
    ft_tf.margin_left = ft_tf.margin_top = ft_tf.margin_right = ft_tf.margin_bottom = 0
    
    fp = ft_tf.paragraphs[0]
    fp.text = "Trường Đại học Công nghệ Thông tin · ĐHQG-TPHCM"
    fp.font.name = "Segoe UI"
    fp.font.size = Pt(8.5)
    fp.font.color.rgb = COLOR_TEXT_FOOTER

def add_bold_bullet(text_frame, bullet_text, font_size=Pt(9.5), space_after=Pt(8)):
    p = text_frame.add_paragraph() if len(text_frame.paragraphs[0].text) > 0 else text_frame.paragraphs[0]
    p.space_after = space_after
    
    # Check if has ":" to bold the leading keyword
    if ":" in bullet_text:
        parts = bullet_text.split(":", 1)
        r_bold = p.add_run()
        r_bold.text = parts[0] + ":"
        r_bold.font.name = "Segoe UI"
        r_bold.font.size = font_size
        r_bold.font.bold = True
        r_bold.font.color.rgb = COLOR_NAVY_DARK
        
        r_text = p.add_run()
        r_text.text = parts[1]
        r_text.font.name = "Segoe UI"
        r_text.font.size = font_size
        r_text.font.bold = False
        r_text.font.color.rgb = COLOR_TEXT_MAIN
    else:
        r = p.add_run()
        r.text = bullet_text
        r.font.name = "Segoe UI"
        r.font.size = font_size
        r.font.color.rgb = COLOR_TEXT_MAIN

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # =========================================================================
    # SLIDE 1: TITLE SLIDE (MATCHING USER'S REFERENCE IMAGE)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    add_branding(s1)
    
    # Tag: B Á O  C Á O  Đ Ề  T À I
    tagBox = s1.shapes.add_textbox(Inches(0.6), Inches(1.35), Inches(6.0), Inches(0.3))
    tf_tag = tagBox.text_frame
    tf_tag.margin_left = tf_tag.margin_top = 0
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = "B Á O   C Á O   Đ Ề   T À I"
    p_tag.font.name = "Segoe UI"
    p_tag.font.size = Pt(11)
    p_tag.font.bold = True
    p_tag.font.color.rgb = COLOR_ORANGE_TAG

    # Main Title
    titleBox = s1.shapes.add_textbox(Inches(0.6), Inches(1.7), Inches(6.5), Inches(1.4))
    tf_t = titleBox.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_top = 0
    pt = tf_t.paragraphs[0]
    pt.text = "Pipeline Phân Tích, Tiền Xử Lý Dữ Liệu IoT & Dự Báo Đa Tầm Vườn Ớt Thông Minh"
    pt.font.name = "Segoe UI"
    pt.font.size = Pt(20)
    pt.font.bold = True
    pt.font.color.rgb = COLOR_NAVY_DARK

    # Subtitle
    subBox = s1.shapes.add_textbox(Inches(0.6), Inches(3.2), Inches(6.5), Inches(0.4))
    tf_sub = subBox.text_frame
    tf_sub.word_wrap = True
    tf_sub.margin_left = tf_sub.margin_top = 0
    ps = tf_sub.paragraphs[0]
    ps.text = "Mục tiêu: Khắc phục lỗi đo cảm biến, lọc bón phân đồng thuận và dự báo 24 giờ tới"
    ps.font.name = "Segoe UI"
    ps.font.size = Pt(10.5)
    ps.font.italic = True
    ps.font.color.rgb = COLOR_UIT_BLUE

    # Subject
    submBox = s1.shapes.add_textbox(Inches(0.6), Inches(3.55), Inches(6.5), Inches(0.3))
    tf_subm = submBox.text_frame
    tf_subm.margin_left = tf_subm.margin_top = 0
    pm = tf_subm.paragraphs[0]
    pm.text = "Môn học: Kỹ thuật Xử lý Dữ liệu (IT2047.26.1.CH.01)"
    pm.font.name = "Segoe UI"
    pm.font.size = Pt(10)
    pm.font.bold = True
    pm.font.color.rgb = COLOR_TEXT_MAIN

    # Card: Giảng viên hướng dẫn
    gv_shape = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(4.0), Inches(4.3), Inches(0.85))
    gv_shape.fill.solid()
    gv_shape.fill.fore_color.rgb = COLOR_CARD_BLUE
    gv_shape.line.color.rgb = RGBColor(219, 234, 254)
    tf_gv = gv_shape.text_frame
    tf_gv.margin_left = Inches(0.2)
    tf_gv.margin_top = Inches(0.12)
    p_gv1 = tf_gv.paragraphs[0]
    p_gv1.text = "GIẢNG VIÊN HƯỚNG DẪN"
    p_gv1.font.name = "Segoe UI"
    p_gv1.font.size = Pt(8.5)
    p_gv1.font.bold = True
    p_gv1.font.color.rgb = COLOR_UIT_BLUE
    p_gv2 = tf_gv.add_paragraph()
    p_gv2.text = "ThS. Võ Phương Bình"
    p_gv2.font.name = "Segoe UI"
    p_gv2.font.size = Pt(13)
    p_gv2.font.bold = True
    p_gv2.font.color.rgb = COLOR_NAVY_DARK

    # Card: Lớp / Khóa
    nhom_shape = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.05), Inches(4.0), Inches(1.85), Inches(0.85))
    nhom_shape.fill.solid()
    nhom_shape.fill.fore_color.rgb = COLOR_NAVY_DARK
    nhom_shape.line.fill.background()
    tf_nhom = nhom_shape.text_frame
    tf_nhom.margin_top = Inches(0.12)
    p_nh1 = tf_nhom.paragraphs[0]
    p_nh1.text = "LỚP CAO HỌC"
    p_nh1.alignment = PP_ALIGN.CENTER
    p_nh1.font.name = "Segoe UI"
    p_nh1.font.size = Pt(8.5)
    p_nh1.font.bold = True
    p_nh1.font.color.rgb = RGBColor(191, 219, 254)
    p_nh2 = tf_nhom.add_paragraph()
    p_nh2.text = "K2025"
    p_nh2.alignment = PP_ALIGN.CENTER
    p_nh2.font.name = "Segoe UI"
    p_nh2.font.size = Pt(14)
    p_nh2.font.bold = True
    p_nh2.font.color.rgb = RGBColor(255, 255, 255)

    # Card: Thành viên nhóm (2 columns)
    tv_shape = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(5.0), Inches(6.3), Inches(1.8))
    tv_shape.fill.solid()
    tv_shape.fill.fore_color.rgb = COLOR_BG_WHITE
    tv_shape.line.color.rgb = RGBColor(226, 232, 240)
    tf_tv = tv_shape.text_frame
    tf_tv.margin_left = Inches(0.2)
    tf_tv.margin_top = Inches(0.12)
    p_tv0 = tf_tv.paragraphs[0]
    p_tv0.text = "THÀNH VIÊN THỰC HIỆN"
    p_tv0.font.name = "Segoe UI"
    p_tv0.font.size = Pt(8.5)
    p_tv0.font.bold = True
    p_tv0.font.color.rgb = COLOR_UIT_BLUE

    # Col 1: Bảo & Khang
    col1_box = s1.shapes.add_textbox(Inches(0.8), Inches(5.35), Inches(2.9), Inches(1.3))
    tf_c1 = col1_box.text_frame
    tf_c1.margin_left = tf_c1.margin_top = 0
    p_b1 = tf_c1.paragraphs[0]
    p_b1.text = "Nguyễn Đình Quốc Bảo\n26848009\n\nNguyễn Lâm Y Khang\n250201065"
    p_b1.font.name = "Segoe UI"
    p_b1.font.size = Pt(9.5)
    p_b1.font.color.rgb = COLOR_TEXT_MAIN

    # Col 2: Khôi & Duy
    col2_box = s1.shapes.add_textbox(Inches(3.8), Inches(5.35), Inches(2.9), Inches(1.3))
    tf_c2 = col2_box.text_frame
    tf_c2.margin_left = tf_c2.margin_top = 0
    p_b2 = tf_c2.paragraphs[0]
    p_b2.text = "Trần Minh Khôi\n250201067\n\nTrương Hoàng Duy\n26848055"
    p_b2.font.name = "Segoe UI"
    p_b2.font.size = Pt(9.5)
    p_b2.font.color.rgb = COLOR_TEXT_MAIN

    # Right Side: Dashboard Mockup
    right_img = os.path.join(IMG_DIR, 'slide_20_demo_web_chili_os.png')
    if os.path.exists(right_img):
        s1.shapes.add_picture(right_img, Inches(7.2), Inches(1.3), width=Inches(5.6))
        
    s1.notes_slide.notes_text_frame.text = (
        "🎬 Transition: Push (Đẩy từ dưới lên - 0.6s)\n"
        "✨ Animation: Fade In thẻ thông tin, Zoom In ảnh Dashboard\n"
        "🎤 Speaker Script: Kính chào Thầy Võ Phương Bình và các bạn. Hôm nay nhóm xin đại diện báo cáo đề tài "
        "'Pipeline Phân tích, Tiền xử lý Dữ liệu IoT & Dự báo Đa Tầm Vườn Ớt Thông Minh'. Đề tài xây dựng một hệ thống hoàn chỉnh từ trạm phần cứng IoT thực tế, "
        "pipeline tiền xử lý khoa học và mô hình AI dự báo 24 giờ."
    )

    # =========================================================================
    # DEFINITIONS FOR SLIDES 02 TO 23
    # =========================================================================
    slides_data = [
        # 02. Agenda
        {
            "num": 2, "title": "Nội Dung Báo Cáo (Agenda)", "sub": "6 giai đoạn thực hiện đề tài",
            "img": None,
            "bullets": [
                "01. Tổng quan & Mục tiêu: Đặt vấn đề canh tác ớt thông minh và thách thức dữ liệu thực tế.",
                "02. Khám phá dữ liệu (EDA): Kiểm toán chuỗi 261 ngày (75.168 mốc) và 5 quy luật nông học then chốt.",
                "03. Kỹ thuật tiền xử lý: Bộ lọc lỗi 3 lớp, bảo vệ đợt bón phân đồng thuận và benchmark điền khuyết.",
                "04. Kiến trúc Pipeline: Quy trình 6 bước chuẩn hóa từ dữ liệu thô Firebase đến dữ liệu sẵn sàng cho AI.",
                "05. Mô hình & Dự báo 24h: So sánh 5 mô hình máy học, chiến lược Direct Knots và dải bất định P10-P90.",
                "06. Web Dashboard & Kết luận: Hệ thống điều hành Chili Farm OS và định hướng mở rộng nghiên cứu."
            ],
            "notes": "🎬 Transition: Push (0.6s)\n✨ Animation: Fade In từng mục Agenda\n🎤 Speaker Script: Bài báo cáo hôm nay gồm 6 nội dung chính, đi từ bối cảnh phần cứng, kiểm toán dữ liệu, các thuật toán tiền xử lý, kiến trúc pipeline cho đến mô hình AI và demo ứng dụng Web."
        },
        # 03. Section Divider 1
        {
            "num": 3, "title": "PHẦN 01 — TỔNG QUAN ĐỀ TÀI", "sub": "Bối cảnh thực tiễn, phát biểu bài toán và mục tiêu giải pháp",
            "img": None,
            "bullets": [
                "Bối cảnh nông nghiệp 4.0: Vườn ớt công nghệ cao đòi hỏi kiểm soát vi khí hậu và dinh dưỡng rễ cây liên tục.",
                "Vấn đề cốt lõi: Cảm biến thực tế hay bị lỗi chết, kẹt và trôi tín hiệu; bón phân qua nước tưới tạo xung vọt dễ bị lọc nhầm.",
                "Mục tiêu nghiên cứu: Chế tạo trạm IoT thực nghiệm, xây dựng pipeline lọc sạch dữ liệu có cơ sở khoa học và dự báo sớm 24 giờ."
            ],
            "notes": "🎬 Transition: Push (0.6s)\n🎤 Speaker Script: Chúng ta bắt đầu với Phần 1: Bối cảnh nghiên cứu và các thách thức lớn khi triển khai trạm IoT nông nghiệp ngoài thực địa."
        },
        # 04. Context, Challenges, Goals
        {
            "num": 4, "title": "Đặt Vấn Đề & Thách Thức Thực Tế", "sub": "Bối cảnh canh tác thông minh, thách thức đo đạc và mục tiêu",
            "img": "station_live_demo.gif",
            "bullets": [
                "Bối cảnh trạm thực địa: Thu thập 10 chỉ số lý hóa đất & vi khí hậu chu kỳ 5 phút/lần bằng bo mạch ESP32.",
                "Thách thức lỗi đo: 814 dòng DHT chết về 0, cảm biến kẹt giá trị >= 60 phút, xung nhọn do nhiễu điện từ.",
                "Thách thức xung đột nông học: Bón phân làm EC & NPK tăng vọt -> Bộ lọc thông thường (IQR/Z-score) xóa nhầm dữ liệu thật.",
                "Mục tiêu giải pháp: Pipeline lọc đa tầng có kiểm chứng và mô hình AI dự báo trực tiếp 24 giờ tới."
            ],
            "notes": "🎬 Transition: Fade (0.5s)\n✨ Animation: Float In 3 thẻ thông tin\n📹 Media: Ảnh động trạm live demo\n🎤 Speaker Script: Khi đặt trạm ngoài vườn ớt, chúng tôi đối mặt 2 thách thức lớn: cảm biến hỏng hóc gây mất dữ liệu và xung đột giữa sự kiện bón phân thật với xung nhiễu kỹ thuật."
        },
        # 05. Section Divider 2
        {
            "num": 5, "title": "PHẦN 02 — TẬP DỮ LIỆU IOT-AGRICULTURE", "sub": "Quy mô chuỗi thời gian 261 ngày, kiểm toán chất lượng và phân tích EDA",
            "img": None,
            "bullets": [
                "Quy mô dữ liệu: 75.168 mốc thời gian liên tục 5 phút từ 14/12/2025 đến 31/08/2026.",
                "Ngoại sinh đồng bộ: Tích hợp 4 biến khí tượng TP.HCM từ Open-Meteo (Bức xạ, Nhiệt độ, Độ ẩm, Mưa).",
                "Phân tích chuyên sâu: 5 biểu đồ EDA trực quan hóa xu hướng chuỗi, chu kỳ ngày/đêm, ma trận tương quan và PCA."
            ],
            "notes": "🎬 Transition: Push (0.6s)\n🎤 Speaker Script: Tiếp theo là Phần 2, chi tiết về tập dữ liệu 75.168 bản ghi và kết quả phân tích khám phá chuyên sâu."
        },
        # 06. Data Scale
        {
            "num": 6, "title": "Quy Mô Dữ Liệu: 75.168 Bản Ghi Chuỗi Thời Gian", "sub": "Độ phủ lưới thời gian 100% — Không thiếu mốc, không trùng lặp",
            "img": "slide_15_eda_monthly_events.png",
            "bullets": [
                "Tổng số mốc thời gian: 75.168 mốc (261 ngày liên tục, bước nhảy chính xác 5 phút).",
                "Độ phủ lưới thời gian: Đạt chuẩn 100.0% (0 bản ghi trùng lặp, 0 bước nhảy bị bỏ sót).",
                "Tích hợp khí tượng: Ghép đồng bộ 4 biến Open-Meteo (Bức xạ mặt trời, Nhiệt độ, Độ ẩm khí tượng, Mưa 5 phút).",
                "Phân chia chuẩn nhân quả: Train (57.312 dòng, ~6.5 tháng) -> Val (Tháng 7, 8.928 dòng) -> Test (Tháng 8, 8.928 dòng)."
            ],
            "notes": "🎬 Transition: Fade (0.5s)\n✨ Animation: Zoom In con số 75.168\n🎤 Speaker Script: Bộ dữ liệu trải dài 261 ngày với 75.168 mốc 5 phút liên tục, đạt độ phủ lưới 100%, không bị mất mốc."
        },
        # 07. Sensors & Hardware
        {
            "num": 7, "title": "Trạm Quan Trắc Thực Địa & Danh Mục Cảm Biến", "sub": "Bo mạch ESP32, Buck XL1509, Cảm biến đất 7-in-1 Modbus RS485",
            "img": "slide_07_iot_station_sensors.png",
            "bullets": [
                "Cảm biến đất 7-in-1 Công nghiệp: Đo Độ ẩm đất (%VWC), Nhiệt đất (°C), EC (mS/cm), pH, N, P, K (mg/kg) qua RS485 Modbus.",
                "Cảm biến vi khí hậu: Nhiệt độ tán lá (°C), Độ ẩm không khí (%RH) qua DHT22, Cường độ ánh sáng (klux).",
                "Khối xử lý & Hiển thị: ESP32 Dual-Core, hạ áp Buck XL1509 ổn định nguồn, màn hình LCD điều khiển tại chỗ.",
                "Đẩy dữ liệu đám mây: Tự động truyền thời gian thực lên Firebase Realtime Database qua Wi-Fi/4G chu kỳ 5 phút."
            ],
            "notes": "🎬 Transition: Fade (0.5s)\n✨ Animation: Wipe From Left các dòng cảm biến\n📹 Media: slide_07_iot_station_sensors.png và station_live_demo.gif\n🎤 Speaker Script: Đây là trạm phần cứng thực tế do nhóm tự chế tạo và lắp đặt. Cảm biến 7-in-1 đo 7 chỉ tiêu đất qua chuẩn Modbus RS485, kết hợp DHT22 và quang trở đo tán lá."
        },
        # 08. Data Features
        {
            "num": 8, "title": "Data Features: 14 Đặc Trưng Dữ Liệu Thu Thập", "sub": "10 chỉ số cảm biến vật lý trạm IoT + 4 chỉ số thời tiết ngoại sinh",
            "img": None,
            "bullets": [
                "Nhóm cảm biến đất: moisture (1-70 %VWC), soiltemp (15-40 °C), pH (3.0-9.5), EC (0.0-4.0 mS/cm), N, P, K (mg/kg).",
                "Nhóm cảm biến tán lá: temperature (5-48 °C), humidity (5-100 %RH), light (0-30 klux).",
                "Nhóm khí tượng ngoại sinh: w_rad (Bức xạ mặt trời W/m²), w_air_T (°C), w_air_RH (%RH), w_rain_5min (mm mưa).",
                "Đặc tính đo lường: Chu kỳ lấy mẫu 5 phút/lần, giá trị đại diện cho trạng thái lý hóa vùng rễ và vi khí hậu vi mô."
            ],
            "notes": "🎬 Transition: Fade (0.5s)\n🎤 Speaker Script: Chúng ta quản lý 14 biến số: 10 biến đo tại chỗ từ trạm và 4 biến thời tiết ngoại sinh bổ trợ cho mô hình AI."
        },
        # 09. Data Coverage
        {
            "num": 9, "title": "Data Coverage: Kiểm Toán Lưới Thời Gian 5 Phút", "sub": "Xác nhận 100% độ phủ thời gian qua 9 tháng ghi nhận",
            "img": None,
            "bullets": [
                "Lưới thời gian lý thuyết: 261 ngày x 288 bước/ngày = 75.168 mốc thời gian.",
                "Lưới thời gian thực tế: 75.168 dòng quan sát -> Missing timestamps = 0, Duplicate timestamps = 0.",
                "Chi tiết theo tháng: Tháng 1, 3, 5, 7, 8 (8.928 dòng/tháng); Tháng 4, 6 (8.640 dòng/tháng); Tháng 2 (8.064 dòng).",
                "Kết luận kiểm toán: Lưới thời gian nguyên vẹn 100%, thuận lợi tuyệt đối cho mô hình học máy chuỗi thời gian."
            ],
            "notes": "🎬 Transition: Fade (0.5s)\n🎤 Speaker Script: Kết quả kiểm toán lưới thời gian cho thấy 0 mốc bị thiếu, 0 mốc trùng lặp. Toàn bộ chuỗi 261 ngày đều có dữ liệu."
        },
        # 10. Data Quality Audit
        {
            "num": 10, "title": "Data Quality: Kiểm Toán Chất Lượng & Bất Thường Vật Lý", "sub": "Nhận diện các lỗi cảm biến vật lý trước khi bước vào tiền xử lý",
            "img": None,
            "bullets": [
                "Lỗi nhiệt độ âm & RH = 0%: Ghi nhận Min temperature = -0.50 °C và RH = 0% -> Lỗi cảm biến DHT mất tín hiệu (TP.HCM không có băng tuyết).",
                "Lỗi quang trở giai đoạn đầu: light có 1.830 ô trống (2.43%) và tương quan yếu do bị tán lá che khuất, đã được nhóm hiệu chuẩn và đưa ra vị trí đo tối ưu.",
                "Đặc tính EC/NPK = 0: Xảy ra chủ yếu khi độ ẩm đất < 12% (đất khô làm điện cực mất tiếp xúc, không phải lỗi phần cứng).",
                "Quy tắc loại bỏ: Thiết lập luật miền vật lý chặn dưới và chặn trên cho từng cảm biến."
            ],
            "notes": "🎬 Transition: Fade (0.5s)\n🎤 Speaker Script: Kiểm toán chất lượng phát hiện lỗi vật lý nghiêm trọng: nhiệt độ âm và độ ẩm 0% khi DHT chết. Đây là cơ sở để nhóm xây dựng bộ lọc vật lý."
        },
        # 11. EDA 1: Daily Trends
        {
            "num": 11, "title": "Data Insights 1: Phân Tích Xu Hướng 261 Ngày", "sub": "Dạng răng cưa dinh dưỡng EC & Nhận diện sự cố bơm hỏng",
            "img": "slide_11_eda_daily_trends.png",
            "bullets": [
                "Độ ẩm đất (45-55%): Duy trì ổn định nhờ tưới tự động; bắt trọn 2 sự cố bơm hỏng (tụt về 24% ngày 22-25/03 & 09-11/08).",
                "Dinh dưỡng EC & NPK: Dạng răng cưa rõ rệt — tăng vọt khi châm phân qua nước tưới, giảm dần do cây hấp thu & rửa trôi.",
                "Mùa mưa (Tháng 5 - 8): Bức xạ giảm, lượng mưa tích lũy cao làm tăng tốc độ rửa trôi ion dinh dưỡng."
            ],
            "notes": "🎬 Transition: Fade (0.5s)\n✨ Animation: Wipe From Left biểu đồ\n🎤 Speaker Script: Biểu đồ xu hướng 261 ngày phản ánh chân thực chu kỳ sinh trưởng của cây ớt: độ ẩm đất được giữ ổn định 45-55%, EC dạng răng cưa sau mỗi lần châm phân."
        },
        # 12. EDA 2: Diurnal Cycles
        {
            "num": 12, "title": "Data Insights 2: Đặc Trưng Chu Kỳ Trong Ngày (Diurnal)", "sub": "Cơ sở lựa chọn đặc trưng sóng tuần hoàn sin/cos(giờ)",
            "img": "slide_12_eda_hourly_cycles.png",
            "bullets": [
                "Nhiệt độ & Bức xạ: Chu kỳ ngày/đêm rất rõ nét, đỉnh nhiệt 13h-14h; nhiệt độ đất có quán tính trễ hơn khí 1.5 - 2 giờ.",
                "Độ ẩm đất & EC: Gần như phẳng theo giờ, biến động chủ yếu theo sự kiện tưới và bón phân.",
                "Quyết định Feature Engineering: Bổ sung biến chu kỳ sin(giờ) và cos(giờ) để mô hình nắm bắt nhịp sinh học ngày/đêm."
            ],
            "notes": "🎬 Transition: Fade (0.5s)\n🎤 Speaker Script: Phân tích 24 giờ cho thấy nhiệt độ không khí và đất có độ lệch pha rõ rệt. Nhóm mã hóa sin và cos theo giờ để mô hình học được chu kỳ ngày đêm."
        },
        # 13. EDA 3: Correlation Matrix
        {
            "num": 13, "title": "Data Insights 3: Ma Trận Tương Quan Đa Biến Pearson", "sub": "Tương quan 10 cảm biến và 4 biến thời tiết ngoại sinh",
            "img": "slide_13_eda_correlation_matrix.png",
            "bullets": [
                "Khối dinh dưỡng trùng lặp: r(EC, K)=0.985, r(EC, P)=0.982, r(EC, N)=0.978 -> Cảm biến 7-in-1 suy diễn từ EC.",
                "Độ ẩm và pH tương quan âm mạnh (r ≈ -0.65): Do tính dẫn điện của dung dịch đất thay đổi khi đất ẩm bão hòa.",
                "Ánh sáng giai đoạn đầu: Tương quan yếu (r=0.02) do quang trở ban đầu bị che khuất, đã được nhóm hiệu chuẩn và đưa ra vị trí đo tối ưu."
            ],
            "notes": "🎬 Transition: Fade (0.5s)\n✨ Animation: Zoom In nhẹ ma trận nhiệt\n🎤 Speaker Script: Ma trận tương quan hé lộ một phát hiện cốt lõi: N, P, K tương quan gần như tuyệt đối với EC (r trên 0.98), cho thấy cảm biến thực chất đo EC rồi tính ra NPK."
        },
        # 14. EDA 4: NPK vs EC & PCA
        {
            "num": 14, "title": "Data Insights 4: Bằng Chứng Thực Nghiệm Về NPK & PCA", "sub": "Phân tích thành phần chính (PCA) giải thích 98.0% phương sai",
            "img": "slide_14_eda_npk_ec_scatter.png",
            "bullets": [
                "Biểu đồ phân tán (Scatter): N, P, K bám sát tuyệt đối vào đường hồi quy tuyến tính của EC.",
                "Phân tích PCA: Thành phần chính đầu tiên (PC1) giải thích đến 98.0% toàn bộ phương sai của cả 4 biến {EC, N, P, K}.",
                "Quyết định giảm chiều: Gom nhóm dinh dưỡng về EC hoặc PC1 giúp giảm triệt để đa cộng tuyến mà không mất thông tin."
            ],
            "notes": "🎬 Transition: Morph (0.6s) hoặc Fade\n🎤 Speaker Script: Phân tích PCA chứng minh PC1 giải thích tới 98.0% phương sai của nhóm NPK và EC. Vì vậy ta có thể rút gọn chiều dữ liệu mà không sợ mất mát thông tin."
        },
        # 15. EDA 5: Train/Val/Test Split
        {
            "num": 15, "title": "Data Insights 5: Chiến Lược Chia Tập Chuỗi Thời Gian", "sub": "Phân chia chuẩn nhân quả theo thời gian — Tránh rò rỉ tương lai",
            "img": "slide_15_eda_monthly_events.png",
            "bullets": [
                "Nguyên tắc bất biến: Tuyệt đối không dùng K-fold xáo trộn ngẫu nhiên (Shuffle) gây rò rỉ dữ liệu tương lai.",
                "Tập Train (14/12/2025 - 30/06/2026): 57.312 bản ghi (~6.5 tháng, chiếm 76.2% tổng dữ liệu).",
                "Tập Validation (Tháng 07/2026): 8.928 bản ghi (chiếm 11.9%, dùng tinh chỉnh siêu tham số mô hình).",
                "Tập Test độc lập (Tháng 08/2026): 8.928 bản ghi (chiếm 11.9%, đánh giá khách quan mùa mưa thực tế)."
            ],
            "notes": "🎬 Transition: Fade (0.5s)\n✨ Animation: Float Up 3 nhãn Train/Val/Test\n🎤 Speaker Script: Chúng tôi kiên quyết tuân thủ nguyên tắc chuỗi thời gian: chia Train, Val và Test theo trục thời gian tuyến tính, nói không với K-fold shuffle."
        },
        # 16. Section Divider 3 & Imputation Benchmark
        {
            "num": 16, "title": "PHẦN 03 — TIỀN XỬ LÝ: BENCHMARK ĐIỀN KHUYẾT", "sub": "Thực nghiệm ẩn ô ngẫu nhiên và đo sai số phục hồi NRMSE",
            "img": "slide_16_preprocessing_imputation.png",
            "bullets": [
                "Đoạn ngắn <= 3 giờ: Nội suy Tuyến tính (Linear) thắng tuyệt đối (NRMSE ≈ 0.10 - 0.17) do tự tương quan 5 phút rất mạnh.",
                "Đoạn dài > 3 giờ trên 1 cảm biến: MICE đa biến vượt trội (NRMSE = 0.3606 so với 0.4403 của Linear) nhờ học tương quan chéo.",
                "Mất kết nối toàn bộ trạm (10 biến cùng mất): Tự động fallback về Tuyến tính vì MICE không còn biến quan sát để mượn thông tin."
            ],
            "notes": "🎬 Transition: Push (0.6s)\n✨ Animation: Wipe Up biểu đồ cột\n🎤 Speaker Script: Trong bài toán điền khuyết, nhóm không đoán mò mà thực nghiệm benchmark: Tuyến tính tối ưu cho đoạn ngắn dưới 3 giờ, còn MICE vượt trội cho đoạn mất dài của 1 cảm biến."
        },
        # 17. Outlier & Fertilizer Consensus
        {
            "num": 17, "title": "Bộ Lọc Lỗi Kết Hợp & Bảo Vệ Đợt Bón Phân", "sub": "Phân biệt giữa sự kiện nông học thật và xung nhiễu cảm biến",
            "img": "slide_17_outlier_event_protection.png",
            "bullets": [
                "Bộ lọc 3 lớp: Luật miền vật lý + Lọc Hampel thích ứng (MAD 24h) + Phát hiện cảm biến kẹt (Stuck >= 12 bước).",
                "Quy tắc đồng thuận nông học: Khi >= 3/4 biến nhóm [EC, N, P, K] cùng tăng vọt -> Xác nhận là BÓN PHÂN THẬT.",
                "Hiệu quả thực nghiệm: Giữ lại 8 điểm của 4 đợt bón phân, giảm 5,1% cờ báo nhầm (156 → 148 dòng)."
            ],
            "notes": "🎬 Transition: Fade (0.5s)\n✨ Animation: Zoom In hộp quy tắc đồng thuận\n🎤 Speaker Script: Đóng góp kỹ thuật quan trọng của nhóm là cơ chế đồng thuận 3 trên 4 biến EC-NPK. Khi phân bón được hòa tan, cả 4 chỉ số cùng nhảy vọt. Quy tắc này cứu sống 4 đợt bón phân không bị bộ lọc outlier xóa oan."
        },
        # 18. Pipeline Architecture
        {
            "num": 18, "title": "Kiến Trúc Pipeline Xử Lý Dữ Liệu 6 Giai Đoạn", "sub": "Quy trình chuẩn hóa đầu-cuối từ dữ liệu thô đến mô hình AI",
            "img": "slide_18_pipeline_architecture.png",
            "bullets": [
                "Giai đoạn 1 - Đọc & Lưới: Flatten Firebase JSON -> Kiểm toán lưới 5 phút (75.168 mốc) -> Ghép Open-Meteo.",
                "Giai đoạn 2 - Bối cảnh vật lý: Lọc 814 dòng DHT chết về 0; giữ EC=0 khi đất khô < 12%; hiệu chuẩn quang trở.",
                "Giai đoạn 3 - Lọc lỗi kết hợp: Miền vật lý + Hampel MAD thích ứng 24h + Phát hiện kẹt tín hiệu (Stuck >= 12).",
                "Giai đoạn 4 - Bảo vệ bón phân: Cơ chế đồng thuận >= 3/4 biến nhóm EC-NPK giữ nguyên đợt bón phân.",
                "Giai đoạn 5 - Điền khuyết & Làm mượt: Hybrid (Tuyến tính đoạn ngắn + MICE đoạn dài) + Làm mượt nhân quả EMA 0.3.",
                "Giai đoạn 6 - Dự báo & Web: Cửa sổ trượt W=12, mô hình Direct Knots 24h kèm dải P10-P90 đẩy lên Chili Farm OS."
            ],
            "notes": "🎬 Transition: Morph (0.6s) hoặc Fade\n✨ Animation: Wipe From Left tuần tự 6 khối\n🎤 Speaker Script: Đây là bức tranh toàn cảnh pipeline 6 giai đoạn của nhóm, khép kín từ khâu tiếp nhận dữ liệu thô, làm sạch, bảo vệ sự kiện, điền khuyết đến dự báo."
        },
        # 19. Model Benchmark & Direct Knots
        {
            "num": 19, "title": "Benchmark Mô Hình Máy Học & Dự Báo Trực Tiếp 24h", "sub": "XGBoost chiến thắng áp đảo với +38.9% Skill & Dải bất định P10-P90",
            "img": "slide_19_model_benchmark.png",
            "bullets": [
                "Tầm h=1 (5 phút): XGBoost (NRMSE=0.1038) và RF (0.1048) cải thiện +14-15% so với Persistence (0.1223).",
                "Tầm h=12 (1 giờ): XGBoost chiến thắng áp đảo (NRMSE=0.2069 so với 0.3387 của Persistence) -> Cải thiện +38.9% Skill.",
                "Đóng góp của Tiền xử lý: Pipeline làm sạch chuẩn khoa học giúp giảm 28.7% sai số dự báo.",
                "Chiến lược Direct Knots 24h: Huấn luyện độc lập từng tầm xa kèm dải bất định P10 - P90, triệt tiêu lỗi tích luỹ đệ quy."
            ],
            "notes": "🎬 Transition: Fade (0.5s)\n✨ Animation: Wipe Up biểu đồ benchmark\n🎤 Speaker Script: Về mặt mô hình, tại tầm 1 giờ, XGBoost đạt NRMSE 0.2069, vượt trội mô hình cơ sở Persistence tới 38.9%. Tiền xử lý chuẩn giúp giảm 28.7% sai số."
        },
        # 20. Demo Web Dashboard
        {
            "num": 20, "title": "Demo Web Dashboard Điều Hành Vườn Ớt (Chili Farm OS)", "sub": "Bản sao kỹ thuật số 3D, Màn hình LCD ảo và Biểu đồ dự báo 24h",
            "img": "slide_20_demo_web_chili_os.png",
            "bullets": [
                "Bản sao kỹ thuật số 3D (Digital Twin 360°): Trực quan hóa trạm quan trắc thực địa sinh động.",
                "Màn hình LCD cảm ứng ảo: Đưa khuyến nghị canh tác tức thời ('Bón vôi nâng pH', 'Tạm dừng bón P-K').",
                "Biểu đồ dự báo 24 giờ: Cảnh báo sớm thiếu ẩm đất trước khi rễ cây tổn thương, điều khiển bơm tưới thông minh.",
                "Nhật ký sự kiện tự động: Lưu vết các đợt bón phân và sự cố thiết bị để phục vụ truy xuất nguồn gốc."
            ],
            "notes": "🎬 Transition: Zoom In (0.6s) hoặc Push\n📹 Media: slide_20_demo_web_chili_os.png hoặc mở trực tiếp web/index.html\n🎤 Speaker Script: Kết quả nghiên cứu được nhóm đóng gói thành Web Dashboard 'Chili Farm OS' với giao diện Digital Twin 3D và khuyến nghị nông học thời gian thực."
        },
        # 21. Conclusions & Future Work
        {
            "num": 21, "title": "Tổng Kết Đề Tài & Hướng Phát Triển Tương Lai", "sub": "Đóng góp khoa học then chốt và lộ trình nâng cấp hệ thống",
            "img": None,
            "bullets": [
                "Đóng góp 1 - Hệ sinh thái hoàn chỉnh: Tự thiết kế trạm IoT phần cứng, pipeline tiền xử lý 6 giai đoạn và Web Dashboard.",
                "Đóng góp 2 - Quy tắc đồng thuận nông học: Giữ lại 8 điểm của 4 đợt bón phân, giảm 5,1% cờ báo nhầm (156 → 148 dòng).",
                "Đóng góp 3 - Hiệu quả dự báo định lượng: Mô hình XGBoost đạt cải thiện +38.9% Skill; tiền xử lý giảm 28.7% sai số.",
                "Hướng phát triển tương lai: Thử nghiệm kiến trúc Deep Learning PatchTST cho chuỗi siêu dài hạn và mở rộng mô hình đa trạm."
            ],
            "notes": "🎬 Transition: Fade (0.5s)\n✨ Animation: Fade In Cột Đóng góp -> Cột Tương lai\n🎤 Speaker Script: Tóm lại, nhóm đã xây dựng thành công trọn gói hệ thống từ trạm IoT phần cứng đến mô hình dự báo và Web OS, đạt hiệu quả cải thiện 38.9% Skill."
        },
        # 22. Q&A & Acknowledgments
        {
            "num": 22, "title": "Hỏi & Đáp (Q&A) — Sẵn Sàng Phản Biện", "sub": "Trân trọng cảm ơn ThS. Võ Phương Bình và Hội đồng phản biện UIT",
            "img": None,
            "bullets": [
                "Nhóm sẵn sàng giải trình các câu hỏi kỹ thuật về:",
                "1. Cơ sở vật lý của việc giữ EC/NPK = 0 khi đất khô dưới 12%.",
                "2. Thuật toán Hampel MAD thích ứng 24h và cơ chế đồng thuận 3/4 biến bón phân.",
                "3. Lý do lựa chọn Linear cho đoạn ngắn <= 3h và MICE cho đoạn dài của 1 cảm biến.",
                "4. Chiến lược Direct Knots kèm dải bất định P10 - P90 và hiệu quả của mô hình XGBoost."
            ],
            "notes": "🎬 Transition: Fade (0.5s)\n🎤 Speaker Script: Nhóm xin chân thành cảm ơn Thầy Võ Phương Bình và các bạn đã lắng nghe. Nhóm sẵn sàng nhận câu hỏi và đóng góp từ Hội đồng."
        },
        # 23. Credits & References
        {
            "num": 23, "title": "Tài Liệu Tham Khảo & Lời Cảm Ơn", "sub": "Trường Đại học Công nghệ Thông tin — Đại học Quốc gia TP.HCM",
            "img": None,
            "bullets": [
                "Hampel, F. R. (1974). The influence curve and its role in robust estimation. JASA.",
                "van Buuren, S. (2011). mice: Multivariate Imputation by Chained Equations in R. JSS.",
                "Chen, T., & Guestrin, C. (2016). XGBoost: A scalable tree boosting system. KDD.",
                "Hyndman, R. J., & Athanasopoulos, G. (2018). Forecasting: principles and practice. OTexts.",
                "Chân thành cảm ơn Khoa Kỹ thuật Máy tính & Bộ môn Hệ thống Thông tin — UIT."
            ],
            "notes": "🎬 Transition: Fade (0.5s)\n🎤 Speaker Script: Slide tài liệu tham khảo chính của đề tài. Xin cảm ơn!"
        }
    ]

    for item in slides_data:
        s = prs.slides.add_slide(blank_layout)
        add_branding(s)
        
        s_num = item["num"]
        s_title = item["title"]
        s_sub = item["sub"]
        img_name = item["img"]
        bullets = item["bullets"]
        notes = item["notes"]
        
        # Slide Title
        t_box = s.shapes.add_textbox(Inches(0.6), Inches(1.15), Inches(12.0), Inches(0.8))
        tf = t_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = 0
        p_t = tf.paragraphs[0]
        p_t.text = f"{s_num:02d}. {s_title}"
        p_t.font.name = "Segoe UI"
        p_t.font.size = Pt(17)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_NAVY_DARK
        
        p_sub = tf.add_paragraph()
        p_sub.text = s_sub
        p_sub.font.name = "Segoe UI"
        p_sub.font.size = Pt(10)
        p_sub.font.italic = True
        p_sub.font.color.rgb = COLOR_UIT_BLUE

        if img_name:
            img_path = os.path.join(IMG_DIR, img_name)
            if os.path.exists(img_path):
                # Place image on right
                s.shapes.add_picture(img_path, Inches(4.5), Inches(2.0), width=Inches(8.2))
                
                # If slide 7, also place small live motion gif at bottom right
                if s_num == 7 and os.path.exists(os.path.join(IMG_DIR, 'station_live_demo.gif')):
                    try:
                        s.shapes.add_picture(os.path.join(IMG_DIR, 'station_live_demo.gif'), Inches(9.8), Inches(4.6), width=Inches(2.8))
                    except Exception as e:
                        print(f"Note: Could not embed gif: {e}")

                # Text notes on left
                n_box = s.shapes.add_textbox(Inches(0.6), Inches(2.0), Inches(3.7), Inches(4.8))
                tf_n = n_box.text_frame
                tf_n.word_wrap = True
                tf_n.margin_left = tf_n.margin_top = 0
                for bullet in bullets:
                    add_bold_bullet(tf_n, bullet, font_size=Pt(9.2), space_after=Pt(8))
        else:
            # Full width text layout
            n_box = s.shapes.add_textbox(Inches(0.6), Inches(2.0), Inches(12.0), Inches(4.8))
            tf_n = n_box.text_frame
            tf_n.word_wrap = True
            tf_n.margin_left = tf_n.margin_top = 0
            for bullet in bullets:
                add_bold_bullet(tf_n, bullet, font_size=Pt(11.5), space_after=Pt(12))

        # Add Presenter Notes
        if notes:
            s.notes_slide.notes_text_frame.text = notes

    prs.save(OUTPUT_PPTX)
    print(f"[SUCCESS] Complete 23-slide PowerPoint successfully created at: {OUTPUT_PPTX}")

if __name__ == '__main__':
    create_presentation()
